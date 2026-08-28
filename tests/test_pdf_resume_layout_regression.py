import json
import tempfile
import unittest
from pathlib import Path

from docx import Document

from core.placeholder_report import build_placeholder_context
from core.renderer import ReportRenderer
from core.resume_parser import parse_resume_for_report


SANITIZED_WRAPPED_PDF_TEXT = """林然
电话：13800000000
邮箱：candidate@example.com
地址：上海
教育背景
工商管理(MBA) 硕士
药物制剂 本科
工作经历
2022.11-至今 齐鲁制药江苏分公司 大区经理
� 管理和制定大区目标 制定年度销售策略、销售目标和指标沟通，制定各地区目标和发展
方向；
� 团队建设 带领团队通过合规和学术化推广，成为业务能力和学术能力最强团队，个人和
地区经理、代表多次获得全国演讲和拜访能力赛前三的奖励。
2018.6- 2020.1 辉瑞制药有限公司 肿瘤及罕见病全国战略执行经理(RSIM)
� 团队领导&人员能力提升：带领全国诊断团队，超额完成全年诊断团队销售
目标，共计搭建超 400 家医院平台；负责全国经理能力提升项目。
2004.01- 2018.06 辉瑞制药有限公司 医学信息沟通地区经理--四星级医学信息沟通地区经理
� 团队领导：带领团队持续五年超 100%高增长达成指标，获得总裁奖 1 人次，SG 最高
荣誉
2 人次，培养和提升 DM8 人
复旦大学
2012.6-2015.6
中国药科大
1996.09-2000.6
"""


class WrappedPdfResumeLayoutRegressionTests(unittest.TestCase):
    def _parsed_and_context(self):
        parsed = parse_resume_for_report(SANITIZED_WRAPPED_PDF_TEXT)
        data = {
            "candidate_name": "林然",
            "position_title": "区域负责人",
            "resume_text": SANITIZED_WRAPPED_PDF_TEXT,
            "original_resume": SANITIZED_WRAPPED_PDF_TEXT,
            "parsed_resume": parsed,
            "recommendation_rationale": {
                "strengths_summary": "候选人具备团队管理与学术推广经验。",
                "risk_notes": "商业化深度需在面试中确认。",
            },
            "motivation": "寻求更具战略纵深的平台。",
            "role_fit": "区域管理经验与目标岗位匹配。",
        }
        context = build_placeholder_context(
            data,
            {"brand_id": "tstar", "brand_name": "T-STAR", "branding": {}},
        )
        return parsed, context, data

    def test_wrapped_pdf_bullets_become_complete_docx_responsibilities(self):
        parsed, context, data = self._parsed_and_context()
        groups_json = json.dumps(context["appendix_blocks"]["experience_groups"], ensure_ascii=False)

        self.assertNotIn("�", parsed["text"])
        self.assertNotIn("�", groups_json)
        self.assertIn("制定各地区目标和发展方向；", groups_json)
        self.assertIn("个人和地区经理、代表多次获得全国演讲和拜访能力赛前三的奖励。", groups_json)
        self.assertIn("共计搭建超 400 家医院平台", groups_json)
        self.assertNotIn('"方向；"', groups_json)
        self.assertNotIn("销售目标，平台", groups_json)

        with tempfile.TemporaryDirectory() as tmp:
            output_path = Path(tmp) / "wrapped-pdf-regression.docx"
            renderer = ReportRenderer({"brand_id": "tstar", "brand_name": "T-STAR", "branding": {}})
            renderer.render(data, output_path)
            document = Document(str(output_path))
            text = "\n".join(paragraph.text for paragraph in document.paragraphs)

        self.assertNotIn("�", text)
        self.assertIn("制定各地区目标和发展方向；", text)
        self.assertNotIn("\n• 方向；", text)

    def test_managerial_title_stays_under_legal_company_name(self):
        _, context, _ = self._parsed_and_context()
        groups = context["appendix_blocks"]["experience_groups"]
        companies = [group["company"] for group in groups]

        self.assertNotIn("肿瘤及罕见病全国战略执行经理(RSIM)", companies)
        pfizer = next(group for group in groups if group["company"] == "辉瑞制药有限公司")
        self.assertEqual(
            [(role["period"], role["title"]) for role in pfizer["roles"]],
            [
                ("2018.6- 2020.1", "肿瘤及罕见病全国战略执行经理(RSIM)"),
                ("2004.01- 2018.06", "医学信息沟通地区经理--四星级医学信息沟通地区经理"),
            ],
        )

    def test_wrapped_honor_and_school_dates_keep_their_sections(self):
        parsed, context, _ = self._parsed_and_context()
        groups_json = json.dumps(context["appendix_blocks"]["experience_groups"], ensure_ascii=False)

        self.assertIn("SG 最高荣誉 2 人次，培养和提升 DM8 人", groups_json)
        self.assertNotIn("2 人次，培养和提升 DM8 人", parsed["structured"]["sections"].get("certificates", []))
        self.assertEqual(
            context["appendix_blocks"]["education"],
            [
                "工商管理(MBA) 硕士",
                "药物制剂 本科",
                "复旦大学 2012.6-2015.6",
                "中国药科大 1996.09-2000.6",
            ],
        )

    def test_tstar_docx_renders_recovered_education(self):
        _, _, data = self._parsed_and_context()

        with tempfile.TemporaryDirectory() as tmp:
            output_path = Path(tmp) / "wrapped-pdf-education.docx"
            renderer = ReportRenderer({"brand_id": "tstar", "brand_name": "T-STAR", "branding": {}})
            renderer.render(data, output_path)
            document = Document(str(output_path))
            text = "\n".join(paragraph.text for paragraph in document.paragraphs)

        self.assertIn("Education / 教育经历", text)
        self.assertIn("复旦大学 2012.6-2015.6", text)
        self.assertIn("中国药科大 1996.09-2000.6", text)


if __name__ == "__main__":
    unittest.main()
