import tempfile
import re
import unittest
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

import docx
from fastapi.testclient import TestClient

from api.bridge import create_app


PROJECT_ROOT = Path(__file__).resolve().parents[1]


def minimal_text_pdf(text: str) -> bytes:
    stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode("ascii")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        b"<< /Length " + str(len(stream)).encode("ascii") + b" >>\nstream\n" + stream + b"\nendstream",
    ]
    chunks = [b"%PDF-1.4\n"]
    offsets = [0]
    for index, obj in enumerate(objects, start=1):
        offsets.append(sum(len(chunk) for chunk in chunks))
        chunks.append(f"{index} 0 obj\n".encode("ascii") + obj + b"\nendobj\n")
    xref_offset = sum(len(chunk) for chunk in chunks)
    chunks.append(f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode("ascii"))
    for offset in offsets[1:]:
        chunks.append(f"{offset:010d} 00000 n \n".encode("ascii"))
    chunks.append(
        f"trailer << /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref_offset}\n%%EOF\n".encode("ascii")
    )
    return b"".join(chunks)


def minimal_docx_xml_text(text: str) -> bytes:
    document_xml = f"""<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
  <w:body>
    <w:p><w:r><w:t>{text}</w:t></w:r></w:p>
  </w:body>
</w:document>
""".encode("utf-8")
    buffer = BytesIO()
    with ZipFile(buffer, "w") as archive:
        archive.writestr("word/document.xml", document_xml)
    return buffer.getvalue()


def minimal_docx_xml_runs_text(runs: list[str]) -> bytes:
    run_xml = "".join(f"<w:r><w:t>{item}</w:t></w:r>" for item in runs)
    document_xml = f"""<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
  <w:body>
    <w:p>{run_xml}</w:p>
  </w:body>
</w:document>
""".encode("utf-8")
    buffer = BytesIO()
    with ZipFile(buffer, "w") as archive:
        archive.writestr("word/document.xml", document_xml)
    return buffer.getvalue()


class ApiBridgeTests(unittest.TestCase):
    def test_health(self):
        with tempfile.TemporaryDirectory() as tmp:
            app = create_app(config_dir=PROJECT_ROOT / "config", data_dir=Path(tmp), public_base_url="http://testserver")
            client = TestClient(app)
            response = client.get("/health")
            self.assertEqual(response.status_code, 200)
            self.assertEqual(response.json()["status"], "ok")

    def test_root_page_and_favicon_are_browser_friendly(self):
        with tempfile.TemporaryDirectory() as tmp:
            app = create_app(config_dir=PROJECT_ROOT / "config", data_dir=Path(tmp), public_base_url="http://testserver")
            client = TestClient(app)

            root_response = client.get("/")
            self.assertEqual(root_response.status_code, 200)
            self.assertIn("Generic Report Tool", root_response.text)
            self.assertIn('name="candidate_name"', root_response.text)
            self.assertIn('name="position_title"', root_response.text)
            self.assertIn('name="salary_info"', root_response.text)
            self.assertIn('name="resume_file"', root_response.text)
            self.assertIn('name="jd_file"', root_response.text)
            self.assertIn('name="jd_text"', root_response.text)
            self.assertIn('name="report_style"', root_response.text)
            self.assertIn('value="tstar" selected', root_response.text)
            self.assertNotIn("Generic Referral Report", root_response.text)
            self.assertIn('value="tstar_warm"', root_response.text)
            self.assertIn('value="consulting_blue"', root_response.text)
            self.assertIn("T-STAR 紫红商务 / 双语模板", root_response.text)
            self.assertNotIn("Logo Yellow", root_response.text)
            self.assertIn('type="file"', root_response.text)
            self.assertIn(".pdf", root_response.text)
            self.assertIn("api/v1/reports/draft-from-files", root_response.text)
            self.assertIn("formatCreateErrorMessage", root_response.text)
            self.assertIn("Upload is too large", root_response.text)
            self.assertIn("上传内容过大", root_response.text)
            self.assertIn("response.status === 413", root_response.text)
            self.assertIn('id="render-btn"', root_response.text)
            self.assertIn('id="render-html-btn"', root_response.text)
            self.assertNotIn('id="render-pdf-btn"', root_response.text)

            favicon_response = client.get("/favicon.ico")
            self.assertEqual(favicon_response.status_code, 204)

    def test_draft_from_files_requires_resume_file(self):
        with tempfile.TemporaryDirectory() as tmp:
            app = create_app(config_dir=PROJECT_ROOT / "config", data_dir=Path(tmp), public_base_url="http://testserver")
            client = TestClient(app)

            response = client.post("/api/v1/reports/draft-from-files", data={
                "candidate_name": "Alice",
                "position_title": "Medical Director",
            })

            self.assertEqual(response.status_code, 422)

    def test_json_draft_rejects_missing_identity_without_resume_reference(self):
        with tempfile.TemporaryDirectory() as tmp:
            app = create_app(config_dir=PROJECT_ROOT / "config", data_dir=Path(tmp), public_base_url="http://testserver")
            client = TestClient(app)

            response = client.post("/api/v1/reports/draft", json={"brand_id": "tstar"})

            self.assertEqual(response.status_code, 422)
            self.assertIn("candidate_name and position_title", response.text)

    def test_draft_from_files_extracts_resume_and_jd(self):
        with tempfile.TemporaryDirectory() as tmp:
            app = create_app(config_dir=PROJECT_ROOT / "config", data_dir=Path(tmp), public_base_url="http://testserver")
            client = TestClient(app)
            resume_file = BytesIO()
            resume_doc = docx.Document()
            resume_doc.add_paragraph("Alice led oncology medical affairs projects.")
            resume_doc.save(resume_file)
            resume_file.seek(0)

            response = client.post(
                "/api/v1/reports/draft-from-files",
                data={
                    "candidate_name": "Alice",
                    "position_title": "Medical Director",
                    "current_company": "Current Pharma",
                    "current_title": "Medical Lead",
                    "salary_info": "Current 80万; Expected 100万",
                    "report_style": "consulting_blue",
                },
                files={
                    "resume_file": (
                        "alice.docx",
                        resume_file.getvalue(),
                        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                    ),
                    "jd_file": (
                        "jd.txt",
                        "Client Company: AstraZeneca\nRole: Medical Director".encode("utf-8"),
                        "text/plain",
                    ),
                },
            )

            self.assertEqual(response.status_code, 200, response.text)
            payload = response.json()
            self.assertTrue(payload["report_id"].startswith("report_"))
            self.assertEqual(payload["source_files"]["resume_file_name"], "alice.docx")
            self.assertEqual(payload["source_files"]["jd_file_name"], "jd.txt")
            self.assertEqual(payload["known_fields"]["client_company"], "AstraZeneca")
            self.assertEqual(payload["known_fields"]["salary_info"], "Current 80万; Expected 100万")
            self.assertIn("oncology medical affairs", payload["data"]["original_resume"])
            self.assertIn("parsed_resume", payload["data"])
            self.assertIn("oncology medical affairs", payload["data"]["parsed_resume"]["text"])
            self.assertEqual(payload["data"]["report_style"], "consulting_blue")

    def test_resume_source_and_candidate_brief_reference_flow(self):
        with tempfile.TemporaryDirectory() as tmp:
            app = create_app(config_dir=PROJECT_ROOT / "config", data_dir=Path(tmp), public_base_url="http://testserver")
            client = TestClient(app)
            resume_text = "\n".join([
                "Name: Alice",
                "Phone: 13800138000",
                "Email: alice@example.com",
                "Professional Experience",
                "AstraZeneca China",
                "2024.01-present Key Account Manager",
                "Led key account access work with measurable growth.",
                "FULL_REFERENCE_FLOW_MARKER",
            ])

            source_response = client.post("/api/v1/resume-sources", json={
                "text": resume_text,
                "file_name": "alice.txt",
            })
            self.assertEqual(source_response.status_code, 200, source_response.text)
            source_payload = source_response.json()
            self.assertRegex(source_payload["resume_source_id"], r"^rs_[a-f0-9]{32}$")
            self.assertGreaterEqual(source_payload["char_count"], len(resume_text))

            brief_response = client.post("/api/v1/candidate-briefs", json={
                "resume_source_id": source_payload["resume_source_id"],
                "known_fields": {
                    "candidate_name": "Alice",
                    "position_title": "KAM",
                    "salary_info": "12K*12 + 20%",
                },
            })
            self.assertEqual(brief_response.status_code, 200, brief_response.text)
            brief_payload = brief_response.json()
            self.assertRegex(brief_payload["candidate_brief_id"], r"^cb_[a-f0-9]{32}$")
            self.assertTrue(brief_payload["candidate_brief"]["career_history"])

            draft_response = client.post("/api/v1/reports/draft", json={
                "brand_id": "tstar",
                "candidate_name": "Alice",
                "position_title": "KAM",
                "candidate_brief_id": brief_payload["candidate_brief_id"],
            })
            self.assertEqual(draft_response.status_code, 200, draft_response.text)
            draft_payload = draft_response.json()
            self.assertIn("FULL_REFERENCE_FLOW_MARKER", draft_payload["data"]["original_resume"])
            self.assertEqual(draft_payload["data"]["candidate_brief_id"], brief_payload["candidate_brief_id"])

    def test_draft_can_be_created_from_candidate_brief_reference_only(self):
        with tempfile.TemporaryDirectory() as tmp:
            app = create_app(config_dir=PROJECT_ROOT / "config", data_dir=Path(tmp), public_base_url="http://testserver")
            client = TestClient(app)
            resume_text = "\n".join([
                "Name: Alice",
                "Phone: 13800138000",
                "Email: alice@example.com",
                "Professional Experience",
                "AstraZeneca China",
                "2024.01-present Key Account Manager",
                "Led key account access work with measurable growth.",
                "FULL_REFERENCE_ONLY_MARKER",
            ])

            source_response = client.post("/api/v1/resume-sources", json={
                "text": resume_text,
                "file_name": "alice.txt",
            })
            self.assertEqual(source_response.status_code, 200, source_response.text)
            brief_response = client.post("/api/v1/candidate-briefs", json={
                "resume_source_id": source_response.json()["resume_source_id"],
                "known_fields": {
                    "candidate_name": "Alice",
                    "position_title": "KAM",
                    "client_company": "Novartis",
                    "salary_info": "12K*12 + 20%",
                },
            })
            self.assertEqual(brief_response.status_code, 200, brief_response.text)

            draft_response = client.post("/api/v1/reports/draft", json={
                "brand_id": "tstar",
                "candidate_brief_id": brief_response.json()["candidate_brief_id"],
            })

            self.assertEqual(draft_response.status_code, 200, draft_response.text)
            payload = draft_response.json()
            self.assertEqual(payload["data"]["candidate_name"], "Alice")
            self.assertEqual(payload["data"]["position_title"], "KAM")
            self.assertEqual(payload["data"]["client_company"], "Novartis")
            self.assertEqual(payload["data"]["salary_info"], "12K*12 + 20%")
            self.assertIn("FULL_REFERENCE_ONLY_MARKER", payload["data"]["original_resume"])

    def test_candidate_brief_rejects_low_quality_resume_source(self):
        with tempfile.TemporaryDirectory() as tmp:
            app = create_app(config_dir=PROJECT_ROOT / "config", data_dir=Path(tmp), public_base_url="http://testserver")
            client = TestClient(app)

            source_response = client.post("/api/v1/resume-sources", json={
                "text": " ".join([
                    "Alice has strong communication skills and excellent stakeholder management capability.",
                    "She is collaborative, resilient, detail oriented, and interested in customer-facing work.",
                    "This text intentionally has no company, role, or employment period evidence.",
                ]),
                "file_name": "low-quality.txt",
            })
            self.assertEqual(source_response.status_code, 200, source_response.text)

            brief_response = client.post("/api/v1/candidate-briefs", json={
                "resume_source_id": source_response.json()["resume_source_id"],
                "known_fields": {
                    "candidate_name": "Alice",
                    "position_title": "KAM",
                },
            })

            self.assertEqual(brief_response.status_code, 422)
            detail = brief_response.json()["detail"]
            self.assertEqual(detail["code"], "resume_quality_blocked")
            self.assertEqual(detail["quality"]["status"], "low_confidence")
            self.assertIn("missing_work_or_project_signal", detail["quality"]["reasons"])
            self.assertEqual(detail["next_action"], "upload_text_resume_or_run_ocr")

    def test_draft_from_files_extracts_docx_xml_text_fallback(self):
        with tempfile.TemporaryDirectory() as tmp:
            app = create_app(config_dir=PROJECT_ROOT / "config", data_dir=Path(tmp), public_base_url="http://testserver")
            client = TestClient(app)
            resume_file = minimal_docx_xml_text("Alice has medical strategy experience in a text box style document.")

            response = client.post(
                "/api/v1/reports/draft-from-files",
                data={
                    "candidate_name": "Alice",
                    "position_title": "Medical Director",
                },
                files={
                    "resume_file": (
                        "alice-xml.docx",
                        resume_file,
                        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                    ),
                },
            )

            self.assertEqual(response.status_code, 200, response.text)
            self.assertIn("medical strategy experience", response.json()["data"]["original_resume"])

    def test_draft_from_files_merges_docx_xml_runs_by_paragraph(self):
        with tempfile.TemporaryDirectory() as tmp:
            app = create_app(config_dir=PROJECT_ROOT / "config", data_dir=Path(tmp), public_base_url="http://testserver")
            client = TestClient(app)
            resume_file = minimal_docx_xml_runs_text(["姓", "名", "：", "杨炯铭"])

            response = client.post(
                "/api/v1/reports/draft-from-files",
                data={
                    "candidate_name": "Yang",
                    "position_title": "地区经理",
                },
                files={
                    "resume_file": (
                        "yang-xml-runs.docx",
                        resume_file,
                        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                    ),
                },
            )

            self.assertEqual(response.status_code, 200, response.text)
            original_resume = response.json()["data"]["original_resume"]
            self.assertIn("姓名：杨炯铭", original_resume)
            self.assertNotIn("姓\n名", original_resume)

    def test_draft_from_files_accepts_jd_text_without_jd_file(self):
        with tempfile.TemporaryDirectory() as tmp:
            app = create_app(config_dir=PROJECT_ROOT / "config", data_dir=Path(tmp), public_base_url="http://testserver")
            client = TestClient(app)

            response = client.post(
                "/api/v1/reports/draft-from-files",
                data={
                    "candidate_name": "Alice",
                    "position_title": "Medical Director",
                    "jd_text": "Client Company: Roche\nLead oncology medical strategy.",
                },
                files={
                    "resume_file": ("alice.txt", b"Alice led oncology launches.", "text/plain"),
                },
            )

            self.assertEqual(response.status_code, 200, response.text)
            payload = response.json()
            self.assertEqual(payload["known_fields"]["client_company"], "Roche")
            self.assertIn("Lead oncology medical strategy", payload["data"]["job_description"])

    def test_draft_from_files_extracts_pdf_resume(self):
        with tempfile.TemporaryDirectory() as tmp:
            app = create_app(config_dir=PROJECT_ROOT / "config", data_dir=Path(tmp), public_base_url="http://testserver")
            client = TestClient(app)
            resume_file = minimal_text_pdf("Alice has oncology launch experience.")

            response = client.post(
                "/api/v1/reports/draft-from-files",
                data={
                    "candidate_name": "Alice",
                    "position_title": "Medical Director",
                },
                files={
                    "resume_file": ("alice.pdf", resume_file, "application/pdf"),
                },
            )

            self.assertEqual(response.status_code, 200, response.text)
            payload = response.json()
            self.assertEqual(payload["source_files"]["resume_file_name"], "alice.pdf")
            self.assertIn("/cards/reports/", payload["card_url"])

    def test_draft_card_and_render_flow(self):
        with tempfile.TemporaryDirectory() as tmp:
            app = create_app(config_dir=PROJECT_ROOT / "config", data_dir=Path(tmp), public_base_url="http://testserver")
            client = TestClient(app)
            draft_response = client.post("/api/v1/reports/draft", json={
                "brand_id": "default",
                "candidate_name": "Alice",
                "position_title": "Medical Director",
            })
            self.assertEqual(draft_response.status_code, 200)
            report_id = draft_response.json()["report_id"]

            card_response = client.get(f"/cards/reports/{report_id}")
            self.assertEqual(card_response.status_code, 200)
            self.assertIn("Alice", card_response.text)
            self.assertIn('id="ai-generate-btn"', card_response.text)
            self.assertIn(f'cards/reports/{report_id}/render', card_response.text)
            self.assertIn(f'cards/reports/{report_id}/render-html', card_response.text)
            self.assertNotIn(f'cards/reports/{report_id}/render-pdf', card_response.text)
            self.assertIn("DOCX Template Diagnostics", card_response.text)

            comments_response = client.post(f"/api/v1/reports/{report_id}/comments", json={})
            self.assertEqual(comments_response.status_code, 200)
            comments_payload = comments_response.json()
            self.assertIn("recommendation_rationale", comments_payload["data"])

            render_response = client.post(f"/api/v1/reports/{report_id}/render")
            self.assertEqual(render_response.status_code, 200)
            self.assertTrue(render_response.json()["filename"].endswith(".docx"))

            browser_render_response = client.post(
                f"/cards/reports/{report_id}/render",
                follow_redirects=False,
            )
            self.assertEqual(browser_render_response.status_code, 303)
            self.assertIn(f"/cards/reports/{report_id}/preview", browser_render_response.headers["location"])
            self.assertIn("format=docx", browser_render_response.headers["location"])

            preview_response = client.get(browser_render_response.headers["location"])
            self.assertEqual(preview_response.status_code, 200)
            self.assertIn("Report Preview", preview_response.text)
            self.assertIn("Download DOCX", preview_response.text)

            html_render_response = client.post(
                f"/cards/reports/{report_id}/render-html",
                follow_redirects=False,
            )
            self.assertEqual(html_render_response.status_code, 303)
            html_preview_response = client.get(html_render_response.headers["location"])
            self.assertEqual(html_preview_response.status_code, 200)
            self.assertIn("<iframe", html_preview_response.text)
            self.assertIn("/preview-files/", html_preview_response.text)
            self.assertIn("Download HTML", html_preview_response.text)
            self.assertIn("Print / Save PDF", html_preview_response.text)
            self.assertIn("预览", html_preview_response.text)

            html_filename = html_render_response.headers["location"].split("filename=", 1)[1]
            inline_response = client.get(f"/preview-files/{html_filename}")
            self.assertEqual(inline_response.status_code, 200)
            self.assertIn("inline", inline_response.headers.get("content-disposition", ""))
            self.assertIn("text/html", inline_response.headers.get("content-type", ""))

    def test_update_report_merges_agent_supplied_fields(self):
        with tempfile.TemporaryDirectory() as tmp:
            app = create_app(config_dir=PROJECT_ROOT / "config", data_dir=Path(tmp), public_base_url="http://testserver")
            client = TestClient(app)
            draft_response = client.post("/api/v1/reports/draft", json={
                "brand_id": "tstar",
                "candidate_name": "Alice",
                "position_title": "Medical Director",
                "resume_text": "Alice led oncology launch work.",
            })
            self.assertEqual(draft_response.status_code, 200, draft_response.text)
            report_id = draft_response.json()["report_id"]

            update_response = client.post(f"/api/v1/reports/{report_id}/update", json={
                "known_fields": {
                    "client_company": "Roche",
                    "salary_info": "Current 80万; expected 100万",
                    "report_style": "consulting_blue",
                    "unknown_field": "should not persist",
                },
                "feedback": "顾问补充了客户和薪资信息。",
            })

            self.assertEqual(update_response.status_code, 200, update_response.text)
            payload = update_response.json()
            self.assertEqual(payload["data"]["client_company"], "Roche")
            self.assertEqual(payload["data"]["salary_info"], "Current 80万; expected 100万")
            self.assertEqual(payload["data"]["report_style"], "consulting_blue")
            self.assertIn("client_company", payload["updated_fields"])
            self.assertIn("unknown_field", payload["ignored_fields"])
            self.assertNotIn("unknown_field", payload["data"])
            self.assertIn("conversation_feedback", payload["data"])
            self.assertIn("/cards/reports/", payload["card_url"])

    def test_comments_can_apply_known_fields_before_ai_generation(self):
        with tempfile.TemporaryDirectory() as tmp:
            app = create_app(config_dir=PROJECT_ROOT / "config", data_dir=Path(tmp), public_base_url="http://testserver")
            client = TestClient(app)
            draft_response = client.post("/api/v1/reports/draft", json={
                "brand_id": "tstar",
                "candidate_name": "Alice",
                "position_title": "Medical Director",
                "resume_text": "Alice led oncology launch work.",
            })
            self.assertEqual(draft_response.status_code, 200, draft_response.text)
            report_id = draft_response.json()["report_id"]

            comments_response = client.post(f"/api/v1/reports/{report_id}/comments", json={
                "known_fields": {"client_company": "Roche"},
                "feedback": "请结合客户要求重新生成。",
            })

            self.assertEqual(comments_response.status_code, 200, comments_response.text)
            self.assertEqual(comments_response.json()["data"]["client_company"], "Roche")

    def test_pdf_report_includes_parsed_resume_across_pages(self):
        from pypdf import PdfReader

        with tempfile.TemporaryDirectory() as tmp:
            app = create_app(config_dir=PROJECT_ROOT / "config", data_dir=Path(tmp), public_base_url="http://testserver")
            client = TestClient(app)
            long_resume = " ".join(
                [f"Experience line {index}: oncology launch achievement {index}." for index in range(45)]
                + ["FINAL_RESUME_MARKER_190_PERCENT"]
            )
            draft_response = client.post("/api/v1/reports/draft", json={
                "brand_id": "tstar",
                "candidate_name": "Alice",
                "position_title": "Medical Director",
                "resume_text": long_resume,
            })
            self.assertEqual(draft_response.status_code, 200, draft_response.text)
            report_id = draft_response.json()["report_id"]

            pdf_response = client.post(f"/api/v1/reports/{report_id}/render-pdf")

            self.assertEqual(pdf_response.status_code, 200, pdf_response.text)
            self.assertRegex(pdf_response.json()["filename"], r"_\d{8}_\d{6}\.pdf$")
            output_path = Path(tmp) / "outputs" / pdf_response.json()["filename"]
            pdf_bytes = output_path.read_bytes()
            self.assertGreater(len(pdf_bytes), 1000)
            reader = PdfReader(str(output_path))
            text = "\n".join(page.extract_text() or "" for page in reader.pages)
            self.assertGreaterEqual(len(reader.pages), 2)
            self.assertIn("Original Resume Appendix", text)
            self.assertIn("FINAL_RESUME_MARKER_190_PERCENT", re.sub(r"\s+", "", text))

    def test_html_report_supports_consulting_blue_style(self):
        with tempfile.TemporaryDirectory() as tmp:
            app = create_app(config_dir=PROJECT_ROOT / "config", data_dir=Path(tmp), public_base_url="http://testserver")
            client = TestClient(app)
            draft_response = client.post("/api/v1/reports/draft", json={
                "brand_id": "tstar",
                "candidate_name": "Alice",
                "position_title": "Medical Director",
                "report_style": "consulting_blue",
                "resume_text": "Alice led oncology launch work. Achieved 120% target in regional medical strategy.",
            })
            self.assertEqual(draft_response.status_code, 200, draft_response.text)
            report_id = draft_response.json()["report_id"]

            html_response = client.post(f"/api/v1/reports/{report_id}/render-html")

            self.assertEqual(html_response.status_code, 200, html_response.text)
            output_path = Path(tmp) / "outputs" / html_response.json()["filename"]
            html_text = output_path.read_text(encoding="utf-8")
            self.assertIn("#0B1F3A", html_text)
            self.assertNotIn("Consulting Blue", html_text)
            self.assertIn("Candidate Profile", html_text)
            self.assertIn("Work Experience", html_text)
            self.assertNotIn("Resume Evidence", html_text)
            self.assertIn("Appendix:", html_text)
            self.assertIn("Original Resume Appendix", html_text)
            self.assertNotIn("Parsing Confidence", html_text)
            self.assertNotIn("Parsed Resume Sections", html_text)
            self.assertNotIn("Structured Resume", html_text)
            self.assertIn("Alice led oncology launch work", html_text)

    def test_html_report_keeps_full_original_resume_appendix(self):
        with tempfile.TemporaryDirectory() as tmp:
            app = create_app(config_dir=PROJECT_ROOT / "config", data_dir=Path(tmp), public_base_url="http://testserver")
            client = TestClient(app)
            resume_text = "\n".join([
                "个人信息",
                "姓名：ZHENG Huang",
                "工作经历",
                "2020.06 - 至今",
                "阿斯利康医药（杭州）有限公司",
                "保留率在65%以上，H1达成93%，泽瑞84%增长。",
                "FULL_ORIGINAL_RESUME_TAIL_MARKER",
            ])
            draft_response = client.post("/api/v1/reports/draft", json={
                "brand_id": "tstar",
                "candidate_name": "ZHENG Huang",
                "position_title": "地区经理",
                "resume_text": resume_text,
            })
            self.assertEqual(draft_response.status_code, 200, draft_response.text)
            report_id = draft_response.json()["report_id"]

            html_response = client.post(f"/api/v1/reports/{report_id}/render-html")

            self.assertEqual(html_response.status_code, 200, html_response.text)
            output_path = Path(tmp) / "outputs" / html_response.json()["filename"]
            html_text = output_path.read_text(encoding="utf-8")
            self.assertIn("Original Resume Appendix", html_text)
            self.assertIn("个人信息", html_text)
            self.assertIn("FULL_ORIGINAL_RESUME_TAIL_MARKER", html_text)
            self.assertIn("company-block", html_text)
            self.assertIn("period", html_text)
            self.assertIn("role-details", html_text)
            self.assertNotIn("Resume Evidence", html_text)
            self.assertNotIn("Template placeholders", html_text)


if __name__ == "__main__":
    unittest.main()
