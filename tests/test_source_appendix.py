import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from docx import Document
from pypdf import PdfWriter


class SourceAppendixTests(unittest.TestCase):
    def test_text_source_preserves_lines_and_blank_lines(self):
        from core.source_appendix import append_source_appendix

        with tempfile.TemporaryDirectory() as tmp:
            source = Path(tmp) / "resume.txt"
            source.write_text("First line\n\nSecond line\n", encoding="utf-8")
            doc = Document()

            append_source_appendix(doc, source)

            values = [paragraph.text for paragraph in doc.paragraphs]
            self.assertIn("First line", values)
            self.assertIn("", values)
            self.assertIn("Second line", values)

    def test_docx_source_is_converted_then_rendered_as_pages(self):
        from core.source_appendix import append_source_appendix

        with tempfile.TemporaryDirectory() as tmp:
            source = Path(tmp) / "resume.docx"
            source.write_bytes(b"docx-placeholder")
            converted = Path(tmp) / "converted.pdf"
            converted.write_bytes(b"pdf-placeholder")
            doc = Document()

            with patch("core.source_appendix.convert_office_to_pdf", return_value=converted) as convert:
                with patch("core.source_appendix.append_pdf_pages") as append_pages:
                    append_source_appendix(doc, source)

            convert.assert_called_once_with(source)
            append_pages.assert_called_once_with(doc, converted)

    def test_structured_only_renderer_has_no_original_resume_appendix(self):
        from core.renderer import ReportRenderer

        renderer = ReportRenderer({"brand_id": "tstar", "brand_name": "T-STAR", "branding": {}})
        doc = renderer._render_tstar_report_v2(
            {
                "candidate_name": "Test Candidate",
                "position_title": "KAM",
                "resume_appendix_mode": "structured_only",
                "parsed_resume": {},
            }
        )

        text = "\n".join(paragraph.text for paragraph in doc.paragraphs)
        self.assertNotIn("Original Resume Appendix", text)

    def test_structured_only_html_has_no_original_resume_appendix(self):
        from core.html_renderer import render_report_html

        html = render_report_html(
            {
                "candidate_name": "Test Candidate",
                "position_title": "KAM",
                "resume_appendix_mode": "structured_only",
                "parsed_resume": {},
            },
            {"brand_id": "tstar", "brand_name": "T-STAR", "branding": {}},
        )

        self.assertNotIn("Original Resume Appendix / 原始简历附录", html)

    def test_structured_only_pdf_fallback_has_no_original_resume_appendix(self):
        from core.pdf_renderer import PdfReportRenderer

        renderer = PdfReportRenderer(
            {"brand_id": "tstar", "brand_name": "T-STAR", "branding": {}}
        )
        renderer._draw_report(
            {
                "candidate_name": "Test Candidate",
                "position_title": "KAM",
                "resume_appendix_mode": "structured_only",
                "original_resume": "This text must not become a fallback appendix.",
                "parsed_resume": {},
            }
        )

        commands = "\n".join(renderer.commands)
        self.assertNotIn("Original Resume Appendix", commands)
        self.assertNotIn("This text must not become a fallback appendix", commands)

    def test_pdf_source_pages_are_embedded_as_images_in_docx(self):
        from core.source_appendix import append_source_appendix

        with tempfile.TemporaryDirectory() as tmp:
            source = Path(tmp) / "resume.pdf"
            writer = PdfWriter()
            writer.add_blank_page(width=595, height=842)
            with source.open("wb") as handle:
                writer.write(handle)
            page = Path(tmp) / "page-1.png"
            page.write_bytes((Path(__file__).parents[1] / "templates" / "tstar" / "tstar_logo_white.png").read_bytes())
            doc = Document()

            with patch("core.source_appendix.render_pdf_pages", return_value=[page]):
                append_source_appendix(doc, source)

            self.assertEqual(len(doc.inline_shapes), 1)
            self.assertTrue(any("Original Resume Appendix" in p.text for p in doc.paragraphs))


if __name__ == "__main__":
    unittest.main()
