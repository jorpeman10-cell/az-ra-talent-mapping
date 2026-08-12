from __future__ import annotations

import shutil
import struct
import subprocess
import tempfile
import base64
import os
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt


TEXT_SUFFIXES = {".txt", ".md"}
OFFICE_SUFFIXES = {".doc", ".docx"}


def append_source_appendix(doc: Document, source_path: str | Path) -> None:
    """Append the preserved source without passing it through resume parsing."""
    source = Path(source_path)
    suffix = source.suffix.lower()
    if suffix in TEXT_SUFFIXES:
        append_text_source(doc, source)
    elif suffix == ".pdf":
        append_pdf_pages(doc, source)
    elif suffix in OFFICE_SUFFIXES:
        converted = convert_office_to_pdf(source)
        try:
            append_pdf_pages(doc, converted)
        finally:
            _cleanup_generated_dir(converted.parent, "resume-source-pdf-")
    else:
        raise ValueError(f"Unsupported source appendix format: {suffix or '(none)'}")


def append_text_source(doc: Document, source_path: Path) -> None:
    doc.add_page_break()
    heading = doc.add_paragraph()
    heading.paragraph_format.space_after = Pt(8)
    run = heading.add_run("Original Resume Appendix / 原始简历附录")
    run.bold = True
    run.font.name = "Microsoft YaHei"
    run.font.size = Pt(14)

    text = source_path.read_text(encoding="utf-8-sig", errors="replace")
    for line in text.splitlines():
        paragraph = doc.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(2)
        paragraph.paragraph_format.line_spacing = 1.15
        run = paragraph.add_run(line if line else " ")
        run.font.name = "Microsoft YaHei"
        run.font.size = Pt(9)


def convert_office_to_pdf(source_path: Path) -> Path:
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        raise RuntimeError("source_appendix_conversion_unavailable: LibreOffice is required")
    output_dir = Path(tempfile.mkdtemp(prefix="resume-source-pdf-"))
    result = subprocess.run(
        [soffice, "--headless", "--convert-to", "pdf", "--outdir", str(output_dir), str(source_path)],
        capture_output=True,
        text=True,
        timeout=120,
        check=False,
    )
    converted = output_dir / f"{source_path.stem}.pdf"
    if result.returncode != 0 or not converted.exists():
        detail = (result.stderr or result.stdout or "conversion failed").strip()
        raise RuntimeError(f"source_appendix_conversion_failed: {detail}")
    return converted


def append_pdf_pages(doc: Document, pdf_path: Path) -> None:
    pages = render_pdf_pages(pdf_path)
    try:
        doc.add_page_break()
        title = doc.add_paragraph()
        title.paragraph_format.space_after = Pt(5)
        run = title.add_run("Original Resume Appendix / 原始简历附录")
        run.bold = True
        run.font.name = "Microsoft YaHei"
        run.font.size = Pt(12)

        section = doc.sections[-1]
        usable_width = section.page_width - section.left_margin - section.right_margin
        usable_height = section.page_height - section.top_margin - section.bottom_margin - Cm(1.2)
        for index, page in enumerate(pages):
            if index:
                doc.add_page_break()
            paragraph = doc.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_after = Pt(0)
            width_px, height_px = _png_dimensions(page)
            ratio = min(usable_width / width_px, usable_height / height_px)
            paragraph.add_run().add_picture(
                str(page), width=int(width_px * ratio), height=int(height_px * ratio)
            )
    finally:
        if pages:
            _cleanup_generated_dir(pages[0].parent, "resume-source-pages-")


def render_pdf_pages(pdf_path: Path) -> list[Path]:
    pdftoppm = shutil.which("pdftoppm")
    if not pdftoppm:
        raise RuntimeError("source_appendix_render_unavailable: pdftoppm is required")
    render_dir = Path(tempfile.mkdtemp(prefix="resume-source-pages-"))
    prefix = render_dir / "page"
    command = [pdftoppm, "-png", "-r", "144", str(pdf_path), str(prefix)]
    if os.name == "nt" and Path(pdftoppm).suffix.lower() in {".cmd", ".bat"}:
        command = ["cmd.exe", "/d", "/c", *command]
    result = subprocess.run(
        command,
        capture_output=True,
        text=True,
        timeout=180,
        check=False,
    )
    pages = sorted(
        render_dir.glob("page-*.png"),
        key=lambda path: int(path.stem.rsplit("-", 1)[-1]),
    )
    if result.returncode != 0 or not pages:
        detail = (result.stderr or result.stdout or "page rendering failed").strip()
        raise RuntimeError(f"source_appendix_render_failed: {detail}")

    return pages


def source_appendix_html(source_path: str | Path) -> str:
    source = Path(source_path)
    suffix = source.suffix.lower()
    if suffix in TEXT_SUFFIXES:
        import html

        lines = source.read_text(encoding="utf-8-sig", errors="replace").splitlines()
        body = "".join(f"<p>{html.escape(line) if line else '&nbsp;'}</p>" for line in lines)
        return f'<section class="page source-text-page"><h2>Original Resume Appendix / 原始简历附录</h2>{body}</section>'
    converted: Path | None = None
    if suffix in OFFICE_SUFFIXES:
        converted = convert_office_to_pdf(source)
        source = converted
    try:
        if source.suffix.lower() != ".pdf":
            raise ValueError(f"Unsupported source appendix format: {suffix or '(none)'}")
        rendered = render_pdf_pages(source)
        try:
            pages = []
            for page in rendered:
                encoded = base64.b64encode(page.read_bytes()).decode("ascii")
                pages.append(
                    '<section class="page source-image-page">'
                    f'<img alt="Original resume page" src="data:image/png;base64,{encoded}">'
                    "</section>"
                )
            return "".join(pages)
        finally:
            if rendered:
                _cleanup_generated_dir(rendered[0].parent, "resume-source-pages-")
    finally:
        if converted is not None:
            _cleanup_generated_dir(converted.parent, "resume-source-pdf-")


def _cleanup_generated_dir(path: Path, prefix: str) -> None:
    if path.name.startswith(prefix):
        shutil.rmtree(path, ignore_errors=True)


def _png_dimensions(path: Path) -> tuple[int, int]:
    with path.open("rb") as handle:
        header = handle.read(24)
    if len(header) != 24 or header[:8] != b"\x89PNG\r\n\x1a\n":
        raise ValueError("Invalid PNG page image")
    return struct.unpack(">II", header[16:24])


__all__ = [
    "append_source_appendix",
    "append_pdf_pages",
    "convert_office_to_pdf",
    "render_pdf_pages",
    "source_appendix_html",
]
